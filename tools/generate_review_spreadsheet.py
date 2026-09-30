#!/usr/bin/env python3
"""
generate_review_spreadsheet.py
Gera a planilha Excel oficial de revisão para SETI: Search for Extraterrestrial Intelligence (Solo Rival).
Abas:
1. Instruções de Revisão
2. Cartas de Ação (S.01 a S.14, S.EXP1)
3. Tiles - Nível I (obj_I_01 a obj_I_04)
4. Tiles - Nível II (obj_II_01 a obj_II_11)
5. Tiles - Nível III (obj_III_01 a obj_III_09)
6. Tiles - Longo Prazo (long_term_01 a long_term_03)
7. Efeitos de Raça (8 espécies alienígenas)
"""

import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Paths
SETI_DIR = "/Users/thiagocarvalho/Documents/Board games/SETI"
BOTS_DIR = "/Users/thiagocarvalho/Documents/Board games/boardbots"
CARDS_JSON = os.path.join(BOTS_DIR, "assets/art/seti/cards_data.json")

OUTPUT_EXCEL_SETI = os.path.join(SETI_DIR, "SETI_Revisao_Rival_Cartas_Tiles_Racas.xlsx")
OUTPUT_EXCEL_BOTS = os.path.join(BOTS_DIR, "tools/SETI_Revisao_Rival_Cartas_Tiles_Racas.xlsx")

# Paletas de cores profissionais
COLOR_HEADER_BG = "1B2A4A"        # Azul Cósmico Escuro
COLOR_HEADER_TEXT = "FFFFFF"      # Branco
COLOR_SUBHEADER_BG = "2C3E50"
COLOR_ZEBRA_LIGHT = "F8FAFC"      # Fundo suave alternado
COLOR_WHITE = "FFFFFF"
COLOR_BORDER = "CBD5E1"           # Cinza claro

# Cores das colunas de revisão (destaque em âmbar suave para o usuário)
COLOR_REV_HEADER_BG = "D97706"    # Âmbar escuro
COLOR_REV_HEADER_TEXT = "FFFFFF"
COLOR_REV_BG = "FEF3C7"           # Amarelo/âmbar bem clarinho
COLOR_REV_BORDER = "F59E0B"

font_title = Font(name="Segoe UI", size=14, bold=True, color="1B2A4A")
font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="475569")
font_header = Font(name="Segoe UI", size=11, bold=True, color=COLOR_HEADER_TEXT)
font_rev_header = Font(name="Segoe UI", size=11, bold=True, color=COLOR_REV_HEADER_TEXT)
font_data = Font(name="Segoe UI", size=10, color="0F172A")
font_data_bold = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
font_rev_data = Font(name="Segoe UI", size=10, color="92400E")

fill_header = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
fill_rev_header = PatternFill(start_color=COLOR_REV_HEADER_BG, end_color=COLOR_REV_HEADER_BG, fill_type="solid")
fill_zebra = PatternFill(start_color=COLOR_ZEBRA_LIGHT, end_color=COLOR_ZEBRA_LIGHT, fill_type="solid")
fill_white = PatternFill(start_color=COLOR_WHITE, end_color=COLOR_WHITE, fill_type="solid")
fill_rev = PatternFill(start_color=COLOR_REV_BG, end_color=COLOR_REV_BG, fill_type="solid")

thin_side = Side(border_style="thin", color=COLOR_BORDER)
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

def apply_styling_and_autofit(ws, start_row=3, rev_cols_start=None):
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = f"A{start_row + 1}"
    
    # Header row height
    ws.row_dimensions[start_row].height = 28
    
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        is_rev_col = rev_cols_start and col[0].column >= rev_cols_start
        
        for cell in col:
            if cell.row < start_row:
                continue
            val = str(cell.value or "")
            lines = val.split("\n")
            line_max = max(len(l) for l in lines) if lines else 0
            if line_max > max_len:
                max_len = line_max
        
        col_width = max(max_len + 4, 12)
        if col_width > 60:
            col_width = 60
        ws.column_dimensions[col_letter].width = col_width

def create_instructions_sheet(wb):
    ws = wb.create_sheet(title="Instruções de Revisão")
    ws.views.sheetView[0].showGridLines = True
    
    ws["B2"] = "Planilha de Revisão Oficial — SETI Solo Rival (Boardbots)"
    ws["B2"].font = Font(name="Segoe UI", size=16, bold=True, color="1B2A4A")
    
    ws["B3"] = "Guia para auditoria e correção das Cartas de Ação, Categorias de Tiles e Efeitos de Raça"
    ws["B3"].font = Font(name="Segoe UI", size=11, italic=True, color="475569")
    
    instructions = [
        ("1. Objetivo desta Planilha:", 
         "Esta planilha consolida todos os dados extraídos das cartas físicas, manuais oficiais (Jogo Base e Expansão Agências Espaciais) e do código atual do bot. Como alguns dados podem conter imprecisões ou interpretações a ajustar, você pode preencher diretamente as colunas destacadas em AMARELO/ÂMBAR."),
        ("2. Estrutura das Abas:",
         "• Cartas de Ação: Lista as 15 cartas de ação do Rival (S.01 a S.14 e S.EXP1) com cada ação (1 a 4), suas setas e prioridades.\n• Tiles - Nível I: 4 objetivos de Nível I (jogo base).\n• Tiles - Nível II: 11 objetivos de Nível II (jogo base e expansão).\n• Tiles - Nível III: 9 objetivos de Nível III (jogo base e expansão).\n• Tiles - Longo Prazo: 3 objetivos de longo prazo da expansão Agências Espaciais.\n• Efeitos de Raça: Todas as 8 espécies alienígenas (5 do jogo base + 3 da expansão) com cartas de ação, regras especiais e pontuação final."),
        ("3. Como Preencher a Revisão (Colunas Amarelas):",
         "• Coluna 'Status': Marque 'OK' para o que estiver certo, ou 'Corrigir' se houver erro.\n• Coluna 'Nova Ordem': Indique a ordem correta da ação (1, 2, 3 ou 4) caso esteja invertida.\n• Coluna 'Correção / Novo Efeito': Escreva o texto ou regra corrigida.\n• Coluna 'Observações': Anote dúvidas ou detalhes de regras para refinarmos no bot."),
        ("4. Próximo Passo:",
         "Após você revisar e salvar a planilha, nós a leremos automaticamente para sincronizar o arquivo 'cards_data.json' e a interface 'bots/seti_bot.html' com 100% de exatidão mecânica.")
    ]
    
    row = 5
    for title, desc in instructions:
        ws.cell(row=row, column=2, value=title).font = Font(name="Segoe UI", size=12, bold=True, color="1B2A4A")
        row += 1
        cell = ws.cell(row=row, column=2, value=desc)
        cell.font = Font(name="Segoe UI", size=10, color="1E293B")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 65 if "\n" in desc else 40
        row += 2
        
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110

def create_cards_sheet(wb, cards_data):
    ws = wb.create_sheet(title="Cartas de Ação")
    
    ws["A1"] = "SETI — Baralho de Ações do Rival (S.01 a S.14 + S.EXP1)"
    ws["A1"].font = font_title
    ws["A2"] = "Revise as ações de cada carta, ordem de execução de cima para baixo e condições de desvio/pulo."
    ws["A2"].font = font_subtitle
    
    headers = [
        "ID Carta", "Tipo", "Seta", "Nº Ação", "Tipo da Ação",
        "Ação / Efeito Atual (PT)", "Ação / Efeito Atual (EN)",
        "Parâmetros & Prioridades", "Condição de Pular Ação",
        "[REVISÃO] Status (OK / Corrigir)", "[REVISÃO] Nova Ordem",
        "[REVISÃO] Ação / Descrição Corrigida (PT)", "[REVISÃO] Observações de Regra"
    ]
    
    start_row = 4
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.alignment = align_header
        if "[REVISÃO]" in h:
            cell.font = font_rev_header
            cell.fill = fill_rev_header
        else:
            cell.font = font_header
            cell.fill = fill_header
        cell.border = border_cell
        
    card_order = [f"S.{i:02d}" for i in range(1, 15)] + ["S.EXP1"]
    current_row = start_row + 1
    
    for card_id in card_order:
        card = cards_data.get(card_id, {})
        card_type = card.get("type", "")
        arrow = "Esquerda (←)" if card.get("arrow") == "left" else "Direita (→)"
        actions = card.get("actions", [])
        
        card_type_pt = "Básica (Inicial)" if card_type == "basic" else ("Avançada" if card_type == "advanced" else "Expansão Inicial")
        
        for act in actions:
            num = act.get("num", 1)
            act_type = act.get("type", "")
            title_pt = act.get("titlePt", "")
            desc_pt = act.get("descPt", "")
            desc_en = act.get("descEn", "")
            
            # Parametros formatados
            params = []
            if "cost" in act:
                params.append(f"Custo: {act['cost']} Pub")
            if "range" in act:
                params.append(f"Alcance: {act['range']}")
            if "planets" in act:
                params.append("Planetas: " + " > ".join(act["planets"]))
            if "preference" in act:
                params.append("Prioridade: " + " > ".join(act["preference"]))
            if "signals" in act:
                params.append(f"Sinais: {act['signals']}")
            params_str = " | ".join(params) if params else "—"
            
            # Condicao de pulo
            skip_cond = "—"
            if act_type == "launch_probe":
                skip_cond = "Pula se o Rival JÁ possuir uma sonda na Terra"
            elif act_type == "analyze":
                skip_cond = "Pula se o Computador do Rival NÃO estiver cheio"
            elif act_type == "tech":
                skip_cond = "Pula se o Rival tiver MENOS de 6 de Publicidade"
            elif act_type == "tech_bonus":
                skip_cond = "Pula se não tiver tecnologia para descartar ou menos de 3 Pub"
            elif act_type == "move_probe":
                skip_cond = "Pula se não houver sonda na Terra ou nenhum planeta alcançável"
            elif act_type == "species_check":
                skip_cond = "Pula se a espécie alienígena correspondente AINDA NÃO foi descoberta"
                
            row_data = [
                card_id, card_type_pt, arrow, num, f"{title_pt} ({act_type})",
                desc_pt, desc_en, params_str, skip_cond,
                "", "", "", "" # Colunas de revisao em branco
            ]
            
            is_zebra = (int(card_id.replace("S.", "").replace("EXP1", "15")) % 2 == 0)
            row_fill = fill_zebra if is_zebra else fill_white
            
            for c_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.border = border_cell
                if c_idx in [1, 2, 3, 4, 10, 11]:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left
                
                if c_idx >= 10: # Colunas de revisao
                    cell.fill = fill_rev
                    cell.font = font_rev_data
                else:
                    cell.fill = row_fill
                    cell.font = font_data_bold if c_idx == 1 else font_data
            
            ws.row_dimensions[current_row].height = 42
            current_row += 1
            
    apply_styling_and_autofit(ws, start_row=4, rev_cols_start=10)

def create_tiles_sheet(wb, category_key, title, subtitle, tile_list):
    ws = wb.create_sheet(title=title)
    
    ws["A1"] = f"SETI — {title}"
    ws["A1"].font = font_title
    ws["A2"] = subtitle
    ws["A2"].font = font_subtitle
    
    headers = [
        "ID do Tile", "Índice", "Arquivo de Imagem", "Categoria",
        "Qtd Tarefas", "Descrição Preliminar das Tarefas / Símbolos",
        "Regra Oficial de Ativação (Manual)",
        "Penalidade Fim de Rodada", "Pontuação Fim de Jogo",
        "[REVISÃO] Status (OK / Corrigir)", "[REVISÃO] Tarefa 1 (Correção)",
        "[REVISÃO] Tarefa 2 (se houver)", "[REVISÃO] Tarefa 3 (se houver)",
        "[REVISÃO] Observações de Regra"
    ]
    
    start_row = 4
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.alignment = align_header
        if "[REVISÃO]" in h:
            cell.font = font_rev_header
            cell.fill = fill_rev_header
        else:
            cell.font = font_header
            cell.fill = fill_header
        cell.border = border_cell
        
    current_row = start_row + 1
    
    for idx, t in enumerate(tile_list, 1):
        tile_id = t["id"]
        img_file = t["file"]
        cat_name = t["category"]
        tasks_count = t.get("tasks", 2)
        desc = t.get("desc", "")
        rules = t.get("rules", "")
        penalty = t.get("penalty", "")
        scoring = t.get("scoring", "")
        
        row_data = [
            tile_id, idx, img_file, cat_name, tasks_count,
            desc, rules, penalty, scoring,
            "", "", "", "", ""
        ]
        
        row_fill = fill_zebra if idx % 2 == 0 else fill_white
        
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=c_idx, value=val)
            cell.border = border_cell
            if c_idx in [1, 2, 4, 5, 10]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
                
            if c_idx >= 10:
                cell.fill = fill_rev
                cell.font = font_rev_data
            else:
                cell.fill = row_fill
                cell.font = font_data_bold if c_idx == 1 else font_data
                
        ws.row_dimensions[current_row].height = 42
        current_row += 1
        
    apply_styling_and_autofit(ws, start_row=4, rev_cols_start=10)

def create_species_sheet(wb, cards_data):
    ws = wb.create_sheet(title="Efeitos de Raça")
    
    ws["A1"] = "SETI — Espécies Alienígenas (Jogo Base & Expansão Agências Espaciais)"
    ws["A1"].font = font_title
    ws["A2"] = "Ações do Rival, ordem, cartas substituídas no baralho e regras específicas de cada espécie."
    ws["A2"].font = font_subtitle
    
    headers = [
        "Espécie / Raça", "Origem", "ID Carta Bot", "Seta", "Substituição no Deck",
        "Nº Ação", "Tipo da Ação", "Descrição da Ação (PT)", "Descrição da Ação (EN)",
        "Regras Solo Específicas (Manual)", "Pontuação Final da Espécie",
        "[REVISÃO] Status (OK / Corrigir)", "[REVISÃO] Nova Ordem",
        "[REVISÃO] Efeito Corrigido (PT)", "[REVISÃO] Observações de Regra"
    ]
    
    start_row = 4
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.alignment = align_header
        if "[REVISÃO]" in h:
            cell.font = font_rev_header
            cell.fill = fill_rev_header
        else:
            cell.font = font_header
            cell.fill = fill_header
        cell.border = border_cell
        
    species_order = [
        ("Mascamitas", "Jogo Base", "species_mascamitas", "Substitui Carta de Espécie 1 ou 2"),
        ("Anomalias", "Jogo Base", "species_anomalias", "Substitui Carta de Espécie 1 ou 2"),
        ("'Oumuamua", "Jogo Base", "species_oumuamua", "Substitui Carta de Espécie 1 ou 2"),
        ("Centaurianos", "Jogo Base", "species_centaurianos", "Substitui Carta de Espécie 1 ou 2"),
        ("Exertianos", "Jogo Base", "species_exertianos", "Substitui Carta de Espécie 1 ou 2"),
        ("Ameba", "Expansão Agências Espaciais", "species_amoeba", "Substitui Carta de Espécie"),
        ("Arkhos", "Expansão Agências Espaciais", "species_arkhos", "Substitui Carta de Espécie"),
        ("Glifídios", "Expansão Agências Espaciais", "species_glyphids", "Substitui Carta de Espécie")
    ]
    
    # Regras especificas extraidas do manual oficial
    species_manual_rules = {
        "Mascamitas": "Amostras coletadas no planeta são colocadas no tabuleiro da espécie. Rival pontua conforme regras normais.",
        "Anomalias": "Se o rival não estiver vencendo a próxima anomalia, marca o vestígio daquela cor e ganha +3 PV.",
        "'Oumuamua": "Quando resolve o ícone de sinal, marca sempre na peça de 'Oumuamua. Coleta exofósseis; só marca espaço com custo se tiver o suficiente (caso contrário considera ocupado).",
        "Centaurianos": "Ao descobrir, rival coloca ficha de marco +15 pontos e pega as restantes. No marco de mensagem, escolhe recompensa mais à esquerda/direita pela seta. Só marca espaço com custo de dados se tiver computador cheio e sobras na reserva.",
        "Exertianos": "Rival só pega e joga cartas Exertianas pela sua carta de ação (ignora marcos). Ao descobrir, ganha avanços por marcos em descoberta. No fim de jogo, todas as cartas jogadas são consideradas concluídas (pontua tudo). Sofre penalidade de perigo (-10% dos pontos totais se tiver mais perigo).",
        "Ameba": "Ao selecionar espaços para marcar no tabuleiro da Ameba, o Rival ignora espaços com 1 ou 0 bônus de organela. As organelas se movem normalmente ao conceder bônus.",
        "Arkhos": "Ao ser descoberto, Rival recebe 3 cartas de segurança. Se ganharia carta de segurança depois, ganha publicidade. A ÚNICA maneira do Rival romper cartas de segurança é através de sua carta especial de ação alienígena. Ignora espaços bloqueados. Ao escolher vestígio com empate mais abaixo, prefere Arkhos.",
        "Glifídios": "Rival ganha fichas de glifo e coloca imediatamente no tabuleiro de tradução (prioridade linha superior; desempate pela seta). Se o espaço superior estiver sem ficha, ignora. No fim de jogo, o Rival marca exatamente 3 PV por ficha de glifo coletada."
    }
    
    species_scoring = {
        "Mascamitas": "Pontuação normal de amostras.",
        "Anomalias": "Pontua anomalias conquistadas.",
        "'Oumuamua": "Pontos normais de setores e exofósseis.",
        "Centaurianos": "+15 pontos a cada ciclo de mensagem completado.",
        "Exertianos": "Totalidade dos pontos de todas as cartas jogadas. Comparação de perigo: quem tiver mais perigo perde 10% dos pontos totais!",
        "Ameba": "Pontuação normal do tabuleiro da Ameba e organelas.",
        "Arkhos": "Pontuação dos níveis acessíveis alcançados na nave.",
        "Glifídios": "3 Pontos de Vitória (PV) por ficha de glifo em posse do Rival (independente do tipo)."
    }
    
    current_row = start_row + 1
    
    for s_idx, (s_name, s_origin, s_id, s_sub) in enumerate(species_order, 1):
        card = cards_data.get(s_id, {})
        arrow = "Esquerda (←)" if card.get("arrow") == "left" else "Direita (→)"
        actions = card.get("actions", [])
        manual_rule = species_manual_rules.get(s_name, "—")
        scoring_rule = species_scoring.get(s_name, "—")
        
        is_zebra = (s_idx % 2 == 0)
        row_fill = fill_zebra if is_zebra else fill_white
        
        for act in actions:
            num = act.get("num", 1)
            act_type = act.get("type", "")
            title_pt = act.get("titlePt", "")
            desc_pt = act.get("descPt", "")
            desc_en = act.get("descEn", "")
            
            row_data = [
                s_name, s_origin, s_id, arrow, s_sub,
                num, f"{title_pt} ({act_type})", desc_pt, desc_en,
                manual_rule, scoring_rule,
                "", "", "", ""
            ]
            
            for c_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.border = border_cell
                if c_idx in [1, 2, 3, 4, 6, 12, 13]:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left
                    
                if c_idx >= 12:
                    cell.fill = fill_rev
                    cell.font = font_rev_data
                else:
                    cell.fill = row_fill
                    cell.font = font_data_bold if c_idx == 1 else font_data
                    
            ws.row_dimensions[current_row].height = 42
            current_row += 1
            
    apply_styling_and_autofit(ws, start_row=4, rev_cols_start=12)

def main():
    print("Iniciando geração da planilha de revisão...")
    
    # Carregar dados existentes
    with open(CARDS_JSON, "r", encoding="utf-8") as f:
        cards_data = json.load(f)
        
    wb = openpyxl.Workbook()
    # Remove aba padrao
    wb.remove(wb.active)
    
    # 1. Instruções
    create_instructions_sheet(wb)
    
    # 2. Cartas de Ação
    create_cards_sheet(wb, cards_data)
    
    # 3. Categorias de Tiles
    # Nivel I (4 tiles)
    tiles_I = [
        {
            "id": f"OBJ-I-{i:02d}", "file": f"obj_I_{i:02d}.webp", "category": "Nível I",
            "tasks": 2,
            "desc": f"Objetivo de Nível I #{i}: 2 tarefas de início de partida (recursos básicos, alcance inicial de sondas ou sinais de rádio).",
            "rules": "Marque cada tarefa assim que o requisito for atingido. Concluído ao marcar ambas.",
            "penalty": "Ao final de R1-R4, se faltar objetivo concluído para descartar, Rival avança +3 na trilha de progresso.",
            "scoring": "Ao final de R5: Rival ganha 5 PV por objetivo restante não concluído na pilha/mesa."
        }
        for i in range(1, 5)
    ]
    create_tiles_sheet(wb, "I", "Tiles - Nível I", "Objetivos de Nível I (Jogo Base — 4 peças)", tiles_I)
    
    # Nivel II (11 tiles)
    tiles_II = [
        {
            "id": f"OBJ-II-{i:02d}", "file": f"obj_II_{i:02d}.webp", "category": "Nível II",
            "tasks": 2,
            "desc": f"Objetivo de Nível II #{i}: Tarefas intermediárias (chegada a planetas com luas, tecnologias específicas, dados em computador, sinais de setores).",
            "rules": "Marque cada tarefa conforme cumprida. Na expansão, objetivos Nível II são revelados desde o setup inicial.",
            "penalty": "Ao final de R1-R4, se faltar objetivo concluído para descartar, Rival avança +3 na trilha de progresso.",
            "scoring": "Ao final de R5: Rival ganha 5 PV por objetivo restante não concluído na pilha/mesa."
        }
        for i in range(1, 12)
    ]
    create_tiles_sheet(wb, "II", "Tiles - Nível II", "Objetivos de Nível II (Jogo Base e Expansão — 11 peças)", tiles_II)
    
    # Nivel III (9 tiles)
    tiles_III = [
        {
            "id": f"OBJ-III-{i:02d}", "file": f"obj_III_{i:02d}.webp", "category": "Nível III",
            "tasks": 2,
            "desc": f"Objetivo de Nível III #{i}: Tarefas avançadas de fim de partida (pouso em planetas externos, grandes volumes de dados/sinais, descoberta de espécies).",
            "rules": "Marque tarefas cumpridas. Na dificuldade 4★ e 5★, mais peças de Nível III entram na pilha.",
            "penalty": "Ao final de R1-R4, se faltar objetivo concluído para descartar, Rival avança +3 na trilha de progresso.",
            "scoring": "Ao final de R5: Rival ganha 5 PV por objetivo restante não concluído na pilha/mesa."
        }
        for i in range(1, 10)
    ]
    create_tiles_sheet(wb, "III", "Tiles - Nível III", "Objetivos de Nível III (Fim de Partida — 9 peças)", tiles_III)
    
    # Longo Prazo (3 tiles da expansao)
    tiles_LT = [
        {
            "id": f"OBJ-LT-{i:02d}", "file": f"long_term_{i:02d}.webp", "category": "Longo Prazo (Expansão)",
            "tasks": 3,
            "desc": f"Objetivo de Longo Prazo #{i} da Expansão Agências Espaciais: Possui exatamente 3 tarefas que exigem esforço cumulativo ao longo da partida.",
            "rules": "Sem reposição: ao concluir as 3 tarefas, remova a peça do jogo (não pode ser usada para cobrir a falta de fim de rodada).",
            "penalty": "Penalidade nas Rodadas 2, 3 e 4: O Rival avança +1 espaço na trilha de progresso para CADA tarefa não concluída ainda visível.",
            "scoring": "Fim de Jogo: NÃO pontua PV para o Rival no final da partida (apenas gera avanço durante o jogo)."
        }
        for i in range(1, 4)
    ]
    create_tiles_sheet(wb, "LT", "Tiles - Longo Prazo", "Objetivos de Longo Prazo da Expansão Agências Espaciais (3 peças)", tiles_LT)
    
    # 4. Efeitos de Raça
    create_species_sheet(wb, cards_data)
    
    # Salvar em ambos os locais
    os.makedirs(os.path.dirname(OUTPUT_EXCEL_SETI), exist_ok=True)
    os.makedirs(os.path.dirname(OUTPUT_EXCEL_BOTS), exist_ok=True)
    
    wb.save(OUTPUT_EXCEL_SETI)
    wb.save(OUTPUT_EXCEL_BOTS)
    print(f"Planilha salva com sucesso em:")
    print(f" -> {OUTPUT_EXCEL_SETI}")
    print(f" -> {OUTPUT_EXCEL_BOTS}")

if __name__ == "__main__":
    main()
